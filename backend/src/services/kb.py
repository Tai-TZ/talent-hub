"""Kho tài liệu cho trợ lý và ban tuyển sinh: nạp tệp, trích văn bản, chia đoạn, tìm kiếm toàn văn (không dấu).

Tệp do người dùng tải lên là dữ liệu không tin cậy: giới hạn kích thước, số trang, số ký tự; trích văn bản bị giới hạn thời gian.
"""

import asyncio
import hashlib
import io
import re
import uuid
from typing import Any

from sqlalchemy import func, select, text, update
from sqlalchemy.sql.elements import ColumnElement

from src.errors import ConflictError, NotFoundError, ValidationFailedError
from src.models import KbChunk, KbDocument
from src.services.audit import RequestMeta, write_audit
from src.services.tenancy import OrgDb

MAX_BYTES = 15 * 1024 * 1024
MAX_CHARS = 2_000_000
MAX_PDF_PAGES = 400
EXTRACT_TIMEOUT_S = 30
CHUNK_CHARS = 900
OVERLAP_CHARS = 150
ALLOWED_EXT = (".txt", ".md", ".markdown", ".csv", ".pdf", ".docx")
HEADING = re.compile(r"^\s{0,3}#{1,4}\s+(.+?)\s*$")
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def _extract_sync(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValidationFailedError(f"Tệp PDF quá dài (tối đa {MAX_PDF_PAGES} trang)")
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return data.decode("utf-8", errors="replace")


async def extract_text(filename: str, data: bytes) -> str:
    if not filename.lower().endswith(ALLOWED_EXT):
        raise ValidationFailedError(
            "Chỉ nhận tệp .txt, .md, .csv, .pdf, .docx", {"file": "Định dạng không được hỗ trợ"}
        )
    if len(data) > MAX_BYTES:
        raise ValidationFailedError("Tệp quá lớn (tối đa 15 MB)", {"file": "Tệp quá lớn"})
    try:
        text_value = await asyncio.wait_for(asyncio.to_thread(_extract_sync, filename, data), EXTRACT_TIMEOUT_S)
    except TimeoutError:
        raise ValidationFailedError(
            "Trích văn bản quá lâu, tệp có thể không hợp lệ", {"file": "Quá thời gian xử lý"}
        ) from None
    except ValidationFailedError:
        raise
    except Exception:  # noqa: BLE001 - tệp hỏng/không đọc được: báo lỗi người dùng, không lộ chi tiết nội bộ
        raise ValidationFailedError("Không đọc được nội dung tệp", {"file": "Tệp hỏng hoặc được bảo vệ"}) from None
    cleaned = CONTROL.sub("", text_value).strip()
    if len(cleaned) < 20:
        raise ValidationFailedError(
            "Tệp không có nội dung văn bản (có thể là ảnh quét)", {"file": "Không trích được văn bản"}
        )
    return cleaned[:MAX_CHARS]


def chunk_text(content: str) -> list[tuple[str, str]]:
    """Chia theo tiêu đề markdown rồi gom đoạn văn đến ~900 ký tự, chồng lấn ~150 ký tự giữa các đoạn liền kề."""
    sections: list[tuple[str, list[str]]] = [("", [])]
    for line in content.splitlines():
        match = HEADING.match(line)
        if match:
            sections.append((match.group(1).strip()[:300], []))
        else:
            sections[-1][1].append(line)

    chunks: list[tuple[str, str]] = []
    for heading, lines in sections:
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", "\n".join(lines)) if p.strip()]
        buffer = ""
        for para in paragraphs:
            pieces = (
                [para[i : i + CHUNK_CHARS] for i in range(0, len(para), CHUNK_CHARS)]
                if len(para) > CHUNK_CHARS
                else [para]
            )
            for piece in pieces:
                if buffer and len(buffer) + len(piece) + 2 > CHUNK_CHARS:
                    chunks.append((heading, buffer))
                    buffer = buffer[-OVERLAP_CHARS:].split(" ", 1)[-1] + "\n\n" + piece
                else:
                    buffer = f"{buffer}\n\n{piece}".strip()
        if buffer:
            chunks.append((heading, buffer))
    return chunks


async def create_document(
    db: OrgDb,
    *,
    actor_user_id: uuid.UUID,
    title: str,
    visibility: str,
    filename: str,
    data: bytes,
    meta: RequestMeta,
) -> KbDocument:
    title = title.strip()
    if not title or len(title) > 200:
        raise ValidationFailedError("Tiêu đề không hợp lệ", {"title": "Nhập tiêu đề (tối đa 200 ký tự)"})
    if visibility not in ("public", "internal"):
        raise ValidationFailedError("Phạm vi không hợp lệ", {"visibility": "Chọn public hoặc internal"})
    content = await extract_text(filename, data)
    checksum = hashlib.sha256(content.encode()).hexdigest()
    dup = (
        await db.session.execute(
            select(KbDocument.title).where(KbDocument.checksum == checksum, KbDocument.status != "retired")
        )
    ).scalar_one_or_none()
    if dup:
        raise ConflictError(f"Nội dung này đã có trong kho: '{dup}'")

    chunks = chunk_text(content)
    doc = KbDocument(
        organization_id=db.org.id,
        title=title,
        visibility=visibility,
        filename=filename[:255],
        content_type="application/octet-stream",
        size_bytes=len(data),
        char_count=len(content),
        chunk_count=len(chunks),
        status="ready",
        checksum=checksum,
        created_by=actor_user_id,
    )
    db.session.add(doc)
    await db.session.flush()
    db.session.add_all(
        KbChunk(organization_id=db.org.id, document_id=doc.id, ordinal=i, heading=heading, content=body)
        for i, (heading, body) in enumerate(chunks)
    )
    write_audit(
        db,
        action="kb.document_added",
        entity_type="kb_document",
        entity_id=doc.id,
        actor_user_id=actor_user_id,
        after={"title": title, "visibility": visibility, "chunks": len(chunks)},
        meta=meta,
    )
    return doc


def doc_out(doc: KbDocument) -> dict[str, Any]:
    return {
        "id": doc.id,
        "title": doc.title,
        "visibility": doc.visibility,
        "filename": doc.filename,
        "size_bytes": doc.size_bytes,
        "char_count": doc.char_count,
        "chunk_count": doc.chunk_count,
        "status": doc.status,
        "version": doc.version,
        "created_at": doc.created_at,
    }


async def list_documents(db: OrgDb, *, status: str | None, limit: int, offset: int) -> tuple[list[KbDocument], int]:
    conditions: list[ColumnElement[bool]] = [KbDocument.status == status] if status else []
    total = (await db.session.execute(select(func.count()).select_from(KbDocument).where(*conditions))).scalar_one()
    rows = (
        (
            await db.session.execute(
                select(KbDocument).where(*conditions).order_by(KbDocument.created_at.desc()).limit(limit).offset(offset)
            )
        )
        .scalars()
        .all()
    )
    return list(rows), total


async def get_document(db: OrgDb, doc_id: uuid.UUID) -> KbDocument:
    doc = (await db.session.execute(select(KbDocument).where(KbDocument.id == doc_id))).scalar_one_or_none()
    if doc is None:
        raise NotFoundError("Không tìm thấy tài liệu")
    return doc


async def preview_chunks(db: OrgDb, doc_id: uuid.UUID, limit: int = 5) -> list[dict[str, Any]]:
    rows = (
        (
            await db.session.execute(
                select(KbChunk).where(KbChunk.document_id == doc_id).order_by(KbChunk.ordinal).limit(limit)
            )
        )
        .scalars()
        .all()
    )
    return [{"ordinal": c.ordinal, "heading": c.heading, "content": c.content[:500]} for c in rows]


async def set_status(
    db: OrgDb, doc_id: uuid.UUID, status: str, *, actor_user_id: uuid.UUID, meta: RequestMeta
) -> KbDocument:
    doc = await get_document(db, doc_id)
    if status == "ready":
        clash = (
            await db.session.execute(
                select(KbDocument.id).where(
                    KbDocument.checksum == doc.checksum, KbDocument.status == "ready", KbDocument.id != doc.id
                )
            )
        ).scalar_one_or_none()
        if clash:
            raise ConflictError("Đã có tài liệu hoạt động với cùng nội dung")
    await db.session.execute(update(KbDocument).where(KbDocument.id == doc.id).values(status=status))
    await db.session.refresh(doc)
    write_audit(
        db,
        action=f"kb.document_{'retired' if status == 'retired' else 'restored'}",
        entity_type="kb_document",
        entity_id=doc.id,
        actor_user_id=actor_user_id,
        meta=meta,
    )
    return doc


_SEARCH_SQL = """
SELECT c.id, c.document_id, d.title, d.visibility, c.heading, c.content, ts_rank_cd(c.tsv, q.query) AS rank
  FROM kb_chunks c
  JOIN kb_documents d ON d.id = c.document_id,
       (SELECT __QUERY__ AS query) q
 WHERE d.status = 'ready' AND c.tsv @@ q.query AND d.visibility = ANY(:visibility)
 ORDER BY rank DESC, c.document_id, c.ordinal
 LIMIT :limit
"""
_AND_SQL = _SEARCH_SQL.replace("__QUERY__", "websearch_to_tsquery('simple', f_unaccent(:q))")
_OR_SQL = _SEARCH_SQL.replace("__QUERY__", "to_tsquery('simple', f_unaccent(:q))")


async def chunks_by_ids(db: OrgDb, chunk_ids: list[uuid.UUID]) -> dict[uuid.UUID, str]:
    """Nội dung đầy đủ của các đoạn (kết quả tìm kiếm chỉ trả đoạn trích ngắn)."""
    if not chunk_ids:
        return {}
    rows = (await db.session.execute(select(KbChunk.id, KbChunk.content).where(KbChunk.id.in_(chunk_ids)))).all()
    return {cid: content for cid, content in rows}


async def search(db: OrgDb, query: str, *, visibility: list[str], limit: int = 5) -> list[dict[str, Any]]:
    """Tìm toàn văn không dấu. Thử khớp mọi từ trước, nếu không có kết quả thì nới lỏng thành khớp bất kỳ từ nào."""
    query = query.strip()[:300]
    if not query:
        return []
    params: dict[str, Any] = {"q": query, "visibility": visibility, "limit": limit}
    rows = (await db.session.execute(text(_AND_SQL), params)).all()
    if not rows:
        tokens = [t for t in re.findall(r"\w+", query, flags=re.UNICODE) if len(t) >= 2][:12]
        if tokens:
            rows = (await db.session.execute(text(_OR_SQL), {**params, "q": " | ".join(tokens)})).all()
    return [
        {
            "chunk_id": r[0],
            "document_id": r[1],
            "title": r[2],
            "visibility": r[3],
            "heading": r[4],
            "snippet": r[5][:400],
            "score": round(float(r[6]), 4),
        }
        for r in rows
    ]
