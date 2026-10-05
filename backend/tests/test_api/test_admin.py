import base64
import io
import uuid
from datetime import UTC, datetime, timedelta

import pytest
from docx import Document
from httpx import AsyncClient
from sqlalchemy import text
from sqlalchemy.ext.asyncio import create_async_engine

from src.config import get_settings
from src.services import email as email_svc
from src.services import jobs
from tests.conftest import OWNER_URL, PASSWORD
from tests.helpers import create_open_intake

GOOD_PASSWORD = "Mật-khẩu-mới-an-toàn-42"


def b64(text_value: str) -> str:
    return base64.b64encode(text_value.encode()).decode()


async def owner_sql(org_id: uuid.UUID, sql: str, params: dict | None = None):  # type: ignore[no-untyped-def]
    engine = create_async_engine(OWNER_URL)
    try:
        async with engine.begin() as conn:
            await conn.execute(text("SELECT set_config('app.org_id', :o, true)"), {"o": str(org_id)})
            result = await conn.execute(text(sql), params or {})
            return result.all() if result.returns_rows else []
    finally:
        await engine.dispose()


async def test_admin_area_requires_permissions(login_as) -> None:  # type: ignore[no-untyped-def]
    paths = ["/admin/users", "/admin/documents", "/admin/costs/summary", "/admin/settings", "/admin/overview"]
    for role in ("applicant", "reviewer", "approver", "training_manager", "mentor"):
        client = await login_as(role)
        for path in paths:
            assert (await client.get(f"/api/v1{path}")).status_code == 403, (role, path)
    admin = await login_as("admin")
    for path in paths:
        assert (await admin.get(f"/api/v1{path}")).status_code == 200, path


async def test_invite_accept_flow_and_single_use(login_as, client: AsyncClient, orgs) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    email = f"new.{uuid.uuid4().hex[:8]}@alpha.test"
    created = await admin.post(
        "/api/v1/admin/users", json={"email": email, "full_name": "Người Mới", "roles": ["reviewer"]}
    )
    assert created.status_code == 201, created.text
    body = created.json()
    assert body["status"] == "invited"
    token = body["invite_link"].rsplit("/", 1)[1]

    # chưa kích hoạt thì không đăng nhập được
    anon = {"X-Organization": "alpha"}
    assert (
        await client.post("/api/v1/auth/login", json={"email": email, "password": GOOD_PASSWORD}, headers=anon)
    ).status_code == 401
    # trùng email trong tổ chức
    assert (
        await admin.post("/api/v1/admin/users", json={"email": email.upper(), "full_name": "X", "roles": ["reviewer"]})
    ).status_code == 409

    preview = await client.get(f"/api/v1/auth/invitations/{token}", headers=anon)
    assert preview.status_code == 200 and preview.json()["kind"] == "invite" and email not in preview.text
    assert (await client.get("/api/v1/auth/invitations/khong-ton-tai-khong-ton-tai", headers=anon)).status_code == 404

    weak = await client.post(
        "/api/v1/auth/invitations/accept", json={"token": token, "password": "short"}, headers=anon
    )
    assert weak.status_code == 400
    ok = await client.post(
        "/api/v1/auth/invitations/accept", json={"token": token, "password": GOOD_PASSWORD}, headers=anon
    )
    assert ok.status_code == 200 and ok.json()["roles"] == ["reviewer"]
    assert (await client.get("/api/v1/me", headers=anon)).status_code == 200  # đã đăng nhập luôn

    # token chỉ dùng một lần
    again = await client.post(
        "/api/v1/auth/invitations/accept", json={"token": token, "password": GOOD_PASSWORD}, headers=anon
    )
    assert again.status_code == 404
    listing = (await admin.get(f"/api/v1/admin/users?q={email}")).json()
    assert listing["items"][0]["status"] == "active"


async def test_expired_invitation_is_rejected(login_as, client: AsyncClient, orgs) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    email = f"exp.{uuid.uuid4().hex[:8]}@alpha.test"
    token = (
        (await admin.post("/api/v1/admin/users", json={"email": email, "full_name": "Hết Hạn", "roles": ["mentor"]}))
        .json()["invite_link"]
        .rsplit("/", 1)[1]
    )
    await owner_sql(
        orgs["alpha"].id, "UPDATE invitations SET expires_at = :t", {"t": datetime.now(UTC) - timedelta(hours=1)}
    )
    res = await client.post(
        "/api/v1/auth/invitations/accept",
        json={"token": token, "password": GOOD_PASSWORD},
        headers={"X-Organization": "alpha"},
    )
    assert res.status_code == 404


async def test_existing_account_cannot_have_password_reset_via_invite(login_as, client: AsyncClient, orgs) -> None:  # type: ignore[no-untyped-def]
    """Admin tổ chức A mời một người đã có tài khoản ở tổ chức B: không thể chiếm tài khoản bằng cách đặt mật khẩu mới."""
    admin = await login_as("admin")
    victim = "mentor@beta.test"  # chỉ thuộc tổ chức beta, có mật khẩu
    created = await admin.post(
        "/api/v1/admin/users", json={"email": victim, "full_name": "Kẻ Tấn Công Đặt Tên", "roles": ["mentor"]}
    )
    assert created.status_code == 201
    token = created.json()["invite_link"].rsplit("/", 1)[1]
    anon = {"X-Organization": "alpha"}
    assert (await client.get(f"/api/v1/auth/invitations/{token}", headers=anon)).json()["has_password"] is True

    hijack = await client.post(
        "/api/v1/auth/invitations/accept", json={"token": token, "password": "Kẻ-tấn-công-đặt-mật-khẩu-1"}, headers=anon
    )
    assert hijack.status_code == 401
    # mật khẩu thật của người đó vẫn dùng được ở tổ chức B
    beta = await client.post(
        "/api/v1/auth/login", json={"email": victim, "password": PASSWORD}, headers={"X-Organization": "beta"}
    )
    assert beta.status_code == 200
    # chủ tài khoản chấp nhận bằng đúng mật khẩu hiện có
    legit = await client.post(
        "/api/v1/auth/invitations/accept", json={"token": token, "password": PASSWORD}, headers=anon
    )
    assert legit.status_code == 200
    # tên không bị admin tổ chức A đổi
    me = (await client.get("/api/v1/me", headers=anon)).json()
    assert me["full_name"] != "Kẻ Tấn Công Đặt Tên"


async def test_email_outbox_is_delivered_by_backend(login_as, orgs) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    email = f"mail.{uuid.uuid4().hex[:8]}@alpha.test"
    await admin.post("/api/v1/admin/users", json={"email": email, "full_name": "Mail Test", "roles": ["applicant"]})
    await jobs.wait_for_all()

    class Capture:
        name = "capture"

        def __init__(self) -> None:
            self.sent: list[tuple[str, str, str]] = []

        async def send(self, to: str, subject: str, body: str) -> None:
            self.sent.append((to, subject, body))

    # email đã được gửi bởi tác vụ nền (console); thêm một email mới rồi gửi bằng backend bắt thư
    await owner_sql(
        orgs["alpha"].id,
        "INSERT INTO email_outbox (id, organization_id, to_email, subject, body) VALUES (gen_random_uuid(), :o, 'x@y.z', 'S', 'B')",
        {"o": orgs["alpha"].id},
    )
    capture = Capture()
    assert await email_svc.flush_outbox(orgs["alpha"].id, capture) >= 1
    assert ("x@y.z", "S", "B") in capture.sent
    rows = await owner_sql(orgs["alpha"].id, "SELECT status FROM email_outbox WHERE to_email = :e", {"e": email})
    assert rows and rows[0][0] == "sent"


async def test_role_changes_and_guards(login_as) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    users = (await admin.get("/api/v1/admin/users?limit=100")).json()
    by_email = {u["email"]: u for u in users["items"]}
    me = by_email["admin@alpha.test"]
    target = by_email["reviewer2@alpha.test"]

    assert "platform_admin" not in users["assignable_roles"]
    assert (
        await admin.patch(f"/api/v1/admin/users/{target['membership_id']}", json={"roles": ["platform_admin"]})
    ).status_code == 422
    assert (await admin.patch(f"/api/v1/admin/users/{target['membership_id']}", json={"roles": []})).status_code == 422
    # không tự đổi vai trò hay khoá chính mình
    assert (
        await admin.patch(f"/api/v1/admin/users/{me['membership_id']}", json={"roles": ["reviewer"]})
    ).status_code == 403
    assert (
        await admin.patch(f"/api/v1/admin/users/{me['membership_id']}", json={"status": "suspended"})
    ).status_code == 403

    changed = await admin.patch(
        f"/api/v1/admin/users/{target['membership_id']}", json={"roles": ["reviewer", "mentor"]}
    )
    assert changed.status_code == 200 and changed.json()["roles"] == ["mentor", "reviewer"]
    await admin.patch(f"/api/v1/admin/users/{target['membership_id']}", json={"roles": ["reviewer"]})


async def test_suspend_blocks_login_and_revokes_sessions(login_as, client: AsyncClient) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    victim = await login_as("applicant3")
    target = next(
        u
        for u in (await admin.get("/api/v1/admin/users?q=applicant3")).json()["items"]
        if u["email"] == "applicant3@alpha.test"
    )
    assert (
        await admin.patch(f"/api/v1/admin/users/{target['membership_id']}", json={"status": "suspended"})
    ).status_code == 200
    assert (await victim.get("/api/v1/me")).status_code == 401  # phiên đang mở bị chặn ngay
    assert (await victim.post("/api/v1/auth/refresh")).status_code == 401
    anon = {"X-Organization": "alpha"}
    assert (
        await client.post(
            "/api/v1/auth/login", json={"email": "applicant3@alpha.test", "password": PASSWORD}, headers=anon
        )
    ).status_code == 401
    assert (
        await admin.patch(f"/api/v1/admin/users/{target['membership_id']}", json={"status": "active"})
    ).status_code == 200
    assert (
        await client.post(
            "/api/v1/auth/login", json={"email": "applicant3@alpha.test", "password": PASSWORD}, headers=anon
        )
    ).status_code == 200


async def test_bulk_import_dry_run_then_apply(login_as) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    tag = uuid.uuid4().hex[:6]
    rows = [
        {"email": f"imp1.{tag}@alpha.test", "full_name": "Một", "roles": ["reviewer"]},
        {"email": f"imp2.{tag}@alpha.test", "full_name": "Hai", "roles": ["mentor", "reviewer"]},
        {"email": "khong-phai-email", "full_name": "Sai", "roles": ["reviewer"]},
        {"email": f"imp1.{tag}@alpha.test", "full_name": "Lặp", "roles": ["reviewer"]},
        {"email": f"imp3.{tag}@alpha.test", "full_name": "Ba", "roles": ["platform_admin"]},
        {"email": "admin@alpha.test", "full_name": "Đã có", "roles": ["admin"]},
    ]
    dry = (await admin.post("/api/v1/admin/users/import", json={"rows": rows, "dry_run": True})).json()
    assert dry["summary"] == {"ok": 2, "created": 0, "error": 4}
    assert (await admin.get(f"/api/v1/admin/users?q=imp1.{tag}")).json()["total"] == 0  # chạy thử không ghi gì
    real = (await admin.post("/api/v1/admin/users/import", json={"rows": rows, "dry_run": False})).json()
    assert real["summary"] == {"ok": 0, "created": 2, "error": 4}
    assert (await admin.get(f"/api/v1/admin/users?q={tag}")).json()["total"] == 2
    assert (await admin.post("/api/v1/admin/users/import", json={"rows": [], "dry_run": True})).status_code == 422


async def test_settings_validation_and_persistence(login_as) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    bad = await admin.put(
        "/api/v1/admin/settings", json={"values": {"usd_vnd_rate": -5, "ai_engine": "gpt", "khong_co": 1}}
    )
    assert bad.status_code == 422 and set(bad.json()["fields"]) == {"usd_vnd_rate", "ai_engine", "khong_co"}
    ok = await admin.put(
        "/api/v1/admin/settings", json={"values": {"usd_vnd_rate": 26000, "ai_monthly_budget_usd": 123}}
    )
    assert ok.status_code == 200
    now = (await admin.get("/api/v1/admin/settings")).json()["values"]
    assert now["usd_vnd_rate"] == 26000 and now["ai_monthly_budget_usd"] == 123 and now["ai_engine"] == "heuristic"
    await admin.put("/api/v1/admin/settings", json={"values": {"usd_vnd_rate": 25500, "ai_monthly_budget_usd": 50}})


async def test_documents_lifecycle_and_accent_insensitive_search(login_as) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    marker = uuid.uuid4().hex[:8]
    content = (
        f"# Quy trình tuyển sinh {marker}\n\nỨng viên nộp hồ sơ trực tuyến, sau đó tham gia bài đánh giá năng lực về tư duy logic và lập trình cơ bản.\n\n"
        f"# Phụ cấp\n\nHọc viên nhận phụ cấp hàng tháng trong suốt thời gian đào tạo tại chương trình {marker}."
    )
    up = await admin.post(
        "/api/v1/admin/documents",
        json={
            "title": f"Cẩm nang {marker}",
            "visibility": "public",
            "filename": "camnang.md",
            "content_base64": b64(content),
        },
    )
    assert up.status_code == 201, up.text
    doc = up.json()
    assert doc["chunk_count"] >= 2 and doc["status"] == "ready"

    # trùng nội dung bị chặn
    dup = await admin.post(
        "/api/v1/admin/documents",
        json={"title": "Bản sao", "visibility": "public", "filename": "b.md", "content_base64": b64(content)},
    )
    assert dup.status_code == 409
    # định dạng lạ, nội dung rỗng, base64 hỏng
    assert (
        await admin.post(
            "/api/v1/admin/documents",
            json={"title": "x", "visibility": "public", "filename": "a.exe", "content_base64": b64("x" * 50)},
        )
    ).status_code == 422
    assert (
        await admin.post(
            "/api/v1/admin/documents",
            json={"title": "x", "visibility": "public", "filename": "a.txt", "content_base64": b64("ngắn")},
        )
    ).status_code == 422
    assert (
        await admin.post(
            "/api/v1/admin/documents",
            json={"title": "x", "visibility": "public", "filename": "a.txt", "content_base64": "###"},
        )
    ).status_code == 422
    assert (
        await admin.post(
            "/api/v1/admin/documents",
            json={"title": "x", "visibility": "secret", "filename": "a.txt", "content_base64": b64("y" * 50)},
        )
    ).status_code == 422

    # tìm không dấu vẫn ra kết quả có dấu
    hits = (await admin.get(f"/api/v1/admin/documents-search?q=phu cap hang thang {marker}")).json()
    assert hits and hits[0]["title"] == f"Cẩm nang {marker}" and "phụ cấp" in hits[0]["snippet"].lower()
    detail = (await admin.get(f"/api/v1/admin/documents/{doc['id']}")).json()
    assert detail["preview"]

    # gỡ khỏi kho thì không tìm thấy nữa, khôi phục thì có lại
    assert (await admin.post(f"/api/v1/admin/documents/{doc['id']}/retire")).json()["status"] == "retired"
    assert (await admin.get(f"/api/v1/admin/documents-search?q={marker}")).json() == []
    assert (await admin.post(f"/api/v1/admin/documents/{doc['id']}/restore")).json()["status"] == "ready"
    assert (await admin.get(f"/api/v1/admin/documents-search?q={marker}")).json()


async def test_docx_upload_is_parsed(login_as) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    marker = uuid.uuid4().hex[:8]
    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph(
        f"Điều kiện tham gia {marker}: sinh viên năm cuối hoặc đã tốt nghiệp các ngành công nghệ thông tin."
    )
    document.add_paragraph("Hồ sơ gồm CV, bảng điểm và các dự án đã thực hiện, nộp trực tuyến trong thời hạn quy định.")
    document.save(buffer)
    up = await admin.post(
        "/api/v1/admin/documents",
        json={
            "title": "Điều kiện",
            "visibility": "internal",
            "filename": "dieukien.docx",
            "content_base64": base64.b64encode(buffer.getvalue()).decode(),
        },
    )
    assert up.status_code == 201, up.text
    assert (await admin.get(f"/api/v1/admin/documents-search?q=dieu kien tham gia {marker}")).json()
    # tệp .docx giả (không phải zip) bị từ chối sạch sẽ
    bogus = await admin.post(
        "/api/v1/admin/documents",
        json={
            "title": "Hỏng",
            "visibility": "public",
            "filename": "hong.docx",
            "content_base64": b64("không phải docx " * 5),
        },
    )
    assert bogus.status_code == 422


async def test_upload_body_limit_is_scoped_to_document_endpoint(login_as) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    big = "A" * (2 * 1024 * 1024)
    assert (
        await admin.post("/api/v1/admin/users/import", content=big, headers={"content-type": "application/json"})
    ).status_code == 413
    ok = await admin.post(
        "/api/v1/admin/documents",
        json={
            "title": "Lớn",
            "visibility": "public",
            "filename": "lon.txt",
            "content_base64": b64("Nội dung dài. " * 150000),
        },
    )
    assert ok.status_code == 201, ok.text


async def test_costs_entries_summary_budget_and_alerts(login_as, orgs) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    today = datetime.now(UTC).date().isoformat()
    tomorrow = (datetime.now(UTC).date() + timedelta(days=1)).isoformat()
    program = (
        await admin.post("/api/v1/programs", json={"code": f"c{uuid.uuid4().hex[:6]}", "name_vi": "Chi phí"})
    ).json()
    cohort = (
        await admin.post(
            f"/api/v1/programs/{program['id']}/cohorts", json={"code": "K1", "name": "Khoá 1", "capacity": 10}
        )
    ).json()

    assert (
        await admin.post(
            "/api/v1/admin/costs/entries", json={"category": "ai", "amount_vnd": 100, "occurred_on": today}
        )
    ).status_code == 422
    assert (
        await admin.post(
            "/api/v1/admin/costs/entries", json={"category": "partner", "amount_vnd": 0, "occurred_on": today}
        )
    ).status_code == 422
    assert (
        await admin.post(
            "/api/v1/admin/costs/entries", json={"category": "partner", "amount_vnd": 100, "occurred_on": tomorrow}
        )
    ).status_code == 422
    assert (
        await admin.post(
            "/api/v1/admin/costs/entries", json={"category": "bogus", "amount_vnd": 100, "occurred_on": today}
        )
    ).status_code == 422

    async def add(category: str, amount: int) -> dict:
        res = await admin.post(
            "/api/v1/admin/costs/entries",
            json={
                "category": category,
                "amount_vnd": amount,
                "occurred_on": today,
                "cohort_id": cohort["id"],
                "description": "test",
            },
        )
        assert res.status_code == 201, res.text
        return res.json()

    stipend = await add("stipend", 8_000_000)
    await add("infrastructure", 2_000_000)
    summary = (await admin.get(f"/api/v1/admin/costs/summary?cohort_id={cohort['id']}")).json()
    assert summary["total"] == 10_000_000 and summary["by_category"]["stipend"]["amount"] == 8_000_000
    assert summary["cost_per_accepted"] is None and summary["alerts"] == []

    # ngân sách: 80% là cảnh báo, 100% là vượt
    assert (
        await admin.put(
            "/api/v1/admin/costs/budgets", json={"cohort_id": cohort["id"], "category": None, "amount_vnd": 11_000_000}
        )
    ).status_code == 200
    warn = (await admin.get(f"/api/v1/admin/costs/summary?cohort_id={cohort['id']}")).json()
    assert warn["overall"]["status"] == "warning" and warn["alerts"][0]["scope"] == "total"
    await admin.put(
        "/api/v1/admin/costs/budgets", json={"cohort_id": cohort["id"], "category": None, "amount_vnd": 9_000_000}
    )
    assert (await admin.get(f"/api/v1/admin/costs/summary?cohort_id={cohort['id']}")).json()["overall"][
        "status"
    ] == "over"

    # huỷ bút toán cần lý do; số liệu cập nhật; không huỷ hai lần
    assert (
        await admin.post(f"/api/v1/admin/costs/entries/{stipend['id']}/void", json={"reason": "ừ"})
    ).status_code == 422
    assert (
        await admin.post(f"/api/v1/admin/costs/entries/{stipend['id']}/void", json={"reason": "Nhập nhầm kỳ"})
    ).status_code == 200
    assert (
        await admin.post(f"/api/v1/admin/costs/entries/{stipend['id']}/void", json={"reason": "Nhập nhầm kỳ"})
    ).status_code == 409
    assert (await admin.get(f"/api/v1/admin/costs/summary?cohort_id={cohort['id']}")).json()["total"] == 2_000_000

    # người không có quyền ghi chi phí
    approver = await login_as("approver")
    assert (
        await approver.post(
            "/api/v1/admin/costs/entries", json={"category": "partner", "amount_vnd": 1, "occurred_on": today}
        )
    ).status_code == 403


async def test_ai_cost_is_derived_from_usage_and_budget_guard_blocks_llm(login_as, app_instance, orgs) -> None:  # type: ignore[no-untyped-def]
    from src.ai.llm_engine import FallbackEngine, LLMEngine
    from src.ai.providers import FakeProvider
    from tests.test_api.test_triage import _prepare

    admin, intake, _ = await _prepare(login_as)
    org_id = orgs["alpha"].id
    before = (await admin.get("/api/v1/admin/costs/summary")).json()
    await owner_sql(
        org_id,
        "INSERT INTO ai_usage (id, organization_id, feature, provider, model, input_tokens, output_tokens, cost_usd) "
        "VALUES (gen_random_uuid(), :o, 'screening', 'anthropic', 'claude-opus-5-5', 1000, 500, 2.5)",
        {"o": org_id},
    )
    after = (await admin.get("/api/v1/admin/costs/summary")).json()
    assert after["ai"]["usd"] == pytest.approx(before["ai"]["usd"] + 2.5)
    assert after["by_category"]["ai"]["amount"] == pytest.approx(after["ai"]["usd"] * after["usd_vnd_rate"], rel=1e-6)
    by_feature = (await admin.get("/api/v1/admin/costs/ai?group=feature")).json()
    assert by_feature[0]["key"] == "screening"

    # chạm trần chi phí AI thì không cho chạy LLM (động cơ offline vẫn chạy)
    await admin.put("/api/v1/admin/settings", json={"values": {"ai_monthly_budget_usd": 0.01}})
    try:
        app_instance.state.screening_engine = FallbackEngine(LLMEngine(FakeProvider(lambda s, u: {"criteria": []})))
        blocked = await admin.post(f"/api/v1/intakes/{intake['id']}/triage", json={})
        assert blocked.status_code == 409 and "trần" in blocked.json()["detail"]
        app_instance.state.screening_engine = None
        ok = await admin.post(f"/api/v1/intakes/{intake['id']}/triage", json={})
        assert ok.status_code == 202  # heuristic không tốn chi phí
        await jobs.wait_for_all()
    finally:
        app_instance.state.screening_engine = None
        await admin.put("/api/v1/admin/settings", json={"values": {"ai_monthly_budget_usd": 50}})


async def test_overview_reports_operational_state(login_as) -> None:  # type: ignore[no-untyped-def]
    admin = await login_as("admin")
    await create_open_intake(admin)
    data = (await admin.get("/api/v1/admin/overview")).json()
    assert data["users"]["by_status"]["active"] >= 5 and "admin" in data["users"]["active_by_role"]
    assert data["intakes"]["open"] >= 1 and data["security_24h"]["logins"] >= 1
    assert data["ai"]["engine"] == "heuristic" and data["email"]["backend"] == get_settings().email_backend
